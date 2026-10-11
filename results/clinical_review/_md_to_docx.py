# -*- coding: utf-8 -*-
"""Markdown -> editable .docx for the clinical-review dossier (Word tables the
reviewer can type into). Handles: #/##/### headings, paragraphs, **bold**/*italic*,
> blockquotes (shaded box), pipe tables (shaded header), '---' rules, '-'/'1.' lists."""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

MD = r"results/clinical_review/dossier_evaluation_clinique.md"
DOCX = r"results/clinical_review/dossier_evaluation_clinique.docx"

NAVY = RGBColor(0x1F, 0x3B, 0x6E)
DARK = RGBColor(0x13, 0x24, 0x2B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEAD_HEX = "1F3B6E"
BOX_HEX = "EEF2F7"
USABLE_CM = 17.0

doc = Document()
# page setup A4 + 2cm margins
sec = doc.sections[0]
sec.page_height = Cm(29.7); sec.page_width = Cm(21.0)
for m in ("top_margin","bottom_margin","left_margin","right_margin"):
    setattr(sec, m, Cm(2.0))
# base font
normal = doc.styles["Normal"]
normal.font.name = "Calibri"; normal.font.size = Pt(10.5); normal.font.color.rgb = DARK

def shade(el, hex_fill):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hex_fill)
    el.append(sh)

def no_borders(tbl):
    tblPr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top","left","bottom","right","insideH","insideV"):
        e = OxmlElement("w:"+edge); e.set(qn("w:val"), "none")
        borders.append(e)
    tblPr.append(borders)

TOKEN = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")
def add_runs(p, text, base_bold=False, color=None, size=None):
    for tok in TOKEN.split(text):
        if not tok:
            continue
        b, it = base_bold, False
        s = tok
        if tok.startswith("**") and tok.endswith("**"):
            b, s = True, tok[2:-2]
        elif tok.startswith("*") and tok.endswith("*"):
            it, s = True, tok[1:-1]
        r = p.add_run(s)
        r.bold = b; r.italic = it
        if color is not None: r.font.color.rgb = color
        if size is not None: r.font.size = Pt(size)

def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]

lines = open(MD, encoding="utf-8").read().splitlines()
i = 0
title_done = False
while i < len(lines):
    raw = lines[i]; s = raw.strip()

    def _clean_head(t):
        return re.sub(r"\*+", "", t)
    if s.startswith("# "):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(_clean_head(s[2:])); r.bold = True; r.font.size = Pt(18); r.font.color.rgb = NAVY
        i += 1; continue
    if s.startswith("## "):
        p = doc.add_paragraph(); p.space_before = Pt(10)
        r = p.add_run(_clean_head(s[3:])); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = NAVY
        i += 1; continue
    if s.startswith("### "):
        p = doc.add_paragraph()
        r = p.add_run(_clean_head(s[4:])); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = DARK
        i += 1; continue

    if s == "---":
        p = doc.add_paragraph(); pPr = p._p.get_or_add_pPr()
        pb = OxmlElement("w:pBdr"); bot = OxmlElement("w:bottom")
        bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "6")
        bot.set(qn("w:space"), "1"); bot.set(qn("w:color"), "C3CED A".replace(" ", ""))
        pb.append(bot); pPr.append(pb)
        i += 1; continue

    if not title_done and s.startswith("*") and s.endswith("*") and "|" not in s:
        while i < len(lines) and lines[i].strip().startswith("*") and lines[i].strip().endswith("*"):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(p, lines[i].strip().strip("*"), color=RGBColor(0x5B,0x6B,0x6E))
            for r in p.runs: r.italic = True; r.font.size = Pt(10.5)
            i += 1
        title_done = True
        continue

    if s.startswith(">"):
        buf = []
        while i < len(lines) and lines[i].strip().startswith(">"):
            buf.append(lines[i].strip()[1:].strip()); i += 1
        text = " ".join(x for x in buf if x)
        tbl = doc.add_table(rows=1, cols=1); tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        no_borders(tbl)
        cell = tbl.cell(0, 0); shade(cell._tc.get_or_add_tcPr(), BOX_HEX)
        cp = cell.paragraphs[0]
        add_runs(cp, text, color=RGBColor(0x24,0x3B,0x55))
        for r in cp.runs: r.italic = True
        doc.add_paragraph()
        continue

    if s.startswith("|") and i+1 < len(lines) and set(lines[i+1].strip()) <= set("|:- "):
        header = split_row(lines[i]); i += 2
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            rows.append(split_row(lines[i])); i += 1
        ncols = len(header)
        tbl = doc.add_table(rows=1, cols=ncols); tbl.style = "Table Grid"
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER; tbl.autofit = False
        # column widths
        if ncols == 4:
            widths = [7.4, 2.3, 3.7, 3.6]
        elif ncols == 2:
            widths = [USABLE_CM/2]*2
        else:
            widths = [USABLE_CM/ncols]*ncols
        hdr = tbl.rows[0].cells
        for c, txt in enumerate(header):
            shade(hdr[c]._tc.get_or_add_tcPr(), HEAD_HEX)
            pp = hdr[c].paragraphs[0]; add_runs(pp, txt, base_bold=True, color=WHITE, size=9.5)
        for r in rows:
            r = (r + [""]*ncols)[:ncols]
            cells = tbl.add_row().cells
            for c, txt in enumerate(r):
                add_runs(cells[c].paragraphs[0], txt, size=9.5)
        for row in tbl.rows:
            for c in range(ncols):
                row.cells[c].width = Cm(widths[c])
        doc.add_paragraph()
        continue

    if re.match(r"^(-|\d+\.)\s+", s):
        m = re.match(r"^(-|\d+\.)\s+(.*)", s)
        prefix = "•  " if m.group(1) == "-" else m.group(1) + "  "
        text = m.group(2); i += 1
        while i < len(lines):
            nx = lines[i].strip()
            if nx == "" or not lines[i].startswith(" ") or nx.startswith(("-","#",">","|")) \
               or re.match(r"^\d+\.\s", nx):
                break
            text += " " + nx; i += 1
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.space_after = Pt(3)
        p.add_run(prefix)
        add_runs(p, text)
        continue

    if s == "":
        i += 1; continue

    buf = [s]; i += 1
    while i < len(lines):
        nx = lines[i].strip()
        if nx == "" or nx.startswith(("#",">","|","---","- ")) or re.match(r"^\d+\.\s", nx):
            break
        buf.append(nx); i += 1
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(5)
    add_runs(p, " ".join(buf))

# footer disclaimer
ft = sec.footer.paragraphs[0]
ft.text = "Prototype de recherche - non destine a un usage clinique en l'etat"
ft.alignment = WD_ALIGN_PARAGRAPH.CENTER
for r in ft.runs: r.font.size = Pt(8); r.font.color.rgb = RGBColor(0x5B,0x6B,0x6E)

doc.save(DOCX)
print("DOCX written:", DOCX)
